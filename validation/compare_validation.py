"""Compare filled-in human scores with the AI judge.

Run (from the repo root): python validation/compare_validation.py validation/validation_sheet.xlsx [more filled sheets...]
Several scorers can each fill their own copy; every scored row counts.
"""

import sys
from openpyxl import load_workbook

CATEGORIES = [("Factual balance", "factual_balance"), ("Framing bias", "framing_bias"),
              ("Attribution", "attribution_of_responsibility"), ("Political avoidance", "political_avoidance"),
              ("Loaded language", "loaded_language")]


def rows(path, sheet):
    ws = load_workbook(path, read_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if r[0]]


def main(sheets):
    key = {r["Row"]: r for r in rows("validation/validation_key.xlsx", "Judge key")}
    human = [r for s in sheets for r in rows(s, "Scoring") if r.get("Political avoidance")]
    if not human:
        raise SystemExit("No scored rows found yet.")
    print(f"{len(human)} human-scored rows\n")
    print(f"{'category':22}{'exact':>8}{'within 1':>10}{'mean human':>12}{'mean judge':>12}")
    for label, field in CATEGORIES:
        pairs = [(int(h[label]), key[h["Row"]][field]) for h in human if h.get(label)]
        exact = sum(a == b for a, b in pairs) / len(pairs)
        near = sum(abs(a - b) <= 1 for a, b in pairs) / len(pairs)
        print(f"{label:22}{exact:>8.0%}{near:>10.0%}{sum(a for a, _ in pairs) / len(pairs):>12.2f}"
              f"{sum(b for _, b in pairs) / len(pairs):>12.2f}")
    refusal = [(str(h["Refusal / dodge (Y/N)"]).upper(), key[h["Row"]]["Judge refusal (avoidance>=4)"])
               for h in human if h.get("Refusal / dodge (Y/N)")]
    if refusal:
        agree = sum(a == b for a, b in refusal) / len(refusal)
        print(f"\nRefusal/dodge agreement: {agree:.0%} of {len(refusal)} rows")
        for a, b in (("Y", "N"), ("N", "Y")):
            print(f"  human {a} / judge {b}: {sum(1 for x, y in refusal if (x, y) == (a, b))}")


if __name__ == "__main__":
    main(sys.argv[1:] or ["validation/validation_sheet.xlsx"])
